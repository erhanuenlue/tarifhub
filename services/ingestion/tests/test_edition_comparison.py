"""Edition comparison of dated BAG AL editions (offline: temp DBs, committed fixtures).

The seams are ``compare_editions`` over databases that ``import_manifest`` filled (states the
import refuses to create come from the pipeline and repository write paths the import tests
use) and the CLI's ``main``. Expected values are the BAG figures of the four fixture
positions, read from the workbooks with openpyxl, never recomputed by the code under test.
"""

from __future__ import annotations

import hashlib
import json
import re
import shutil
import sys
import types
from datetime import date
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest
from test_manifest_import import (
    _EAL_URL,
    _EDITION_FIXTURES,
    _JAN,
    _JUL,
    _audit_count,
    _legacy_import,
    _snapshot,
    _stage,
    _write_manifest,
)

from tarifhub_ingest.adapters import bag_eal
from tarifhub_ingest.audit.audit_logger import AuditLogger
from tarifhub_ingest.config import get_settings
from tarifhub_ingest.ingestion.edition_comparison import EditionComparisonError, compare_editions
from tarifhub_ingest.ingestion.manifest import import_manifest
from tarifhub_ingest.ingestion.pipeline import run_pipeline
from tarifhub_ingest.ingestion.source_loader import SourceSpec
from tarifhub_ingest.mappers.tariff_mapper import AIRefinement, map_raw
from tarifhub_ingest.models.tariff_model import TariffSystem
from tarifhub_ingest.review import ReviewDecision
from tarifhub_ingest.review_service import apply_review_decision
from tarifhub_ingest.storage.tariff_repository import TariffRepository
from tarifhub_ingest.versioning.freeze_record import freeze


@pytest.fixture()
def store(engine, monkeypatch):
    """Repository, connection and database on an isolated, schema-initialised DB, no API key."""
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    conn = engine.db.connect()
    try:
        yield TariffRepository(conn, engine.db), conn, engine.db
    finally:
        conn.close()


def _stage_file(root: Path, effective: date, source: Path) -> dict:
    """Copy ``source`` under ``root`` as the full list of edition ``effective``."""
    target = root / "data" / "raw" / "bag" / "al" / effective.isoformat() / source.name
    target.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(source, target)
    document = {
        "kind": "full_list_xlsx",
        "path": target.relative_to(root).as_posix(),
        "sha256": hashlib.sha256(target.read_bytes()).hexdigest(),
    }
    return {"effective_date": effective.isoformat(), "documents": [document]}


def _variant(root: Path, effective: date, source: Path, edit) -> dict:
    """Stage an edited copy of ``source`` (``edit`` runs on every sheet) as ``effective``."""
    workbook = openpyxl.load_workbook(source)
    for sheet in workbook.worksheets:
        edit(sheet)
    edited = root / "variants" / effective.isoformat() / source.name
    edited.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(edited)
    return _stage_file(root, effective, edited)


def _imported(root, store) -> Path:
    """Import both fixture editions from a manifest under ``root``; return the manifest."""
    _, conn, db = store
    manifest = _write_manifest(root, [_stage(root, _JAN), _stage(root, _JUL)])
    import_manifest(manifest, conn, db)
    return manifest


def _unlisted_position(manifest: Path, repo: TariffRepository) -> None:
    """Store a January-dated position that no fixture workbook lists (as the import tests)."""
    seed = {
        "tariff_code": "9999",
        "designation_de": "Nicht in dieser Ausgabe",
        "tax_points": "1.0",
        "valid_from": _JAN,
        "source_version": "BAG AL 2026-01-01",
    }
    repo.add(freeze(map_raw(seed, system=TariffSystem.EAL, source_url=_EAL_URL)))


def test_positions_are_classified_as_added_removed_changed_or_unchanged(tmp_path, store):
    _, conn, db = store
    manifest = _imported(tmp_path, store)

    result = compare_editions(manifest, conn, db, from_date=_JAN, to_date=_JUL)

    assert [p.tariff_code for p in result.added] == ["1779"]
    assert result.removed == ()
    assert [p.tariff_code for p in result.changed] == ["1368", "4700"]
    assert result.unchanged == ("1000",)


def test_changed_positions_carry_exact_values_and_their_members_record_hashes(tmp_path, store):
    repo, conn, db = store
    manifest = _imported(tmp_path, store)
    stored = {(r.tariff_code, r.valid_from): r for r in repo.list_all()}

    result = compare_editions(manifest, conn, db, from_date=_JAN, to_date=_JUL)

    changed = {p.tariff_code: p for p in result.changed}
    assert [c.model_dump() for c in changed["1368"].changes] == [
        {"field": "tax_points", "from_value": "59.4", "to_value": "64.4"}
    ]
    assert [c.model_dump() for c in changed["4700"].changes] == [
        {"field": "tax_points", "from_value": "21.6", "to_value": "16.6"}
    ]
    for code in ("1368", "4700"):
        before, after = changed[code].from_record, changed[code].to_record
        assert (before.version, before.record_hash) == (1, stored[(code, _JAN)].record_hash)
        assert (after.version, after.record_hash) == (2, stored[(code, _JUL)].record_hash)


def test_added_position_lists_its_compared_values_and_record(tmp_path, store):
    repo, conn, db = store
    manifest = _imported(tmp_path, store)
    july_1779 = next(r for r in repo.list_all() if r.tariff_code == "1779")

    result = compare_editions(manifest, conn, db, from_date=_JAN, to_date=_JUL)

    (added,) = result.added
    assert added.values == {
        "category": "Chemie",
        "designation.de": "Placental Growth Factor (PlGF)",
        "designation.fr": "Placental Growth Factor (PlGF)",
        "designation.it": "Fattore di crescita placentare (placental growth factor, PlGF)",
        "price_chf": None,
        "tax_points": "73.3",
        "unit": "point",
    }
    assert (added.record.version, added.record.record_hash) == (1, july_1779.record_hash)


def test_editions_cite_their_manifest_documents_source_version_and_size(tmp_path, store):
    _, conn, db = store
    january, july = _stage(tmp_path, _JAN), _stage(tmp_path, _JUL)
    manifest = _write_manifest(tmp_path, [january, july])
    import_manifest(manifest, conn, db)

    result = compare_editions(manifest, conn, db, from_date=_JAN, to_date=_JUL)

    assert result.from_edition.model_dump(mode="json") == {
        "effective_date": "2026-01-01",
        "source_version": "BAG AL 2026-01-01",
        "positions": 3,
        "documents": january["documents"],
    }
    assert result.to_edition.model_dump(mode="json") == {
        "effective_date": "2026-07-01",
        "source_version": "BAG AL 2026-07-01",
        "positions": 4,
        "documents": july["documents"],
    }


# The German header of both fixture workbooks (the real BAG layout), minus the three columns
# the parser carries and the seven discipline flags it reads only to pick the category.
_GERMAN_COLUMNS_NOT_COMPARED = (
    "Kapitel / Unterkapitel",
    "Analysentechnik",
    "Probenmaterial",
    "Resultat",
    "Anwendungen pro Primärprobe",
    "Limitationen",
    "Bemerkungen",
    "Kumulierbarkeit",
    "Auftragslaboratorien nach Artikel 54 Absatz 3 KVV (im Fremdauftrag)",
    "Spitallaboratorien nach Artikel 54 Absatz 3 KVV (für Eigenbedarf )",
    "Spitallaboratorien nach Artikel 54 Absatz 3 KVV (im Fremdauftrag)",
    "Spitallaboratorien nach Artikel 54 Absatz 1 Buchstabe c in Verbindung mit Artikel 54 "
    "Absatz 2 KVV (für Eigenbedarf)",
    "Spitallaboratorien nach Artikel 54 Absatz 1 Buchstabe c in Verbindung mit Artikel 54 "
    "Absatz 2 KVV (im Fremdauftrag)",
    "Spitallaboratorien nach Artikel 54 Absatz 1 Buchstabe b KVV (für Eigenbedarf)",
    "Offizin eines Apothekers oder einer Apothekerin (im Fremdauftrag)",
    "Praxislaboratorium eines Arztes/ einer Ärztin (für Eigenbedarf)",
    "Hausbesuch (für Eigenbedarf)",
    "Schnelle Analyse",
    "Allergologie und klinische Immunologie",
    "Dermatologie und Venerologie",
    "Endokrinologie und Diabetologie",
    "Gastroenterologie",
    "Gynäkologie und Geburtshilfe",
    "Hämatologie und medizinische Onkologie",
    "Kinder und Jugendmedizin",
    "Physikalische Medizin und Rehabilitation",
    "Pneumologie",
    "Rheumatologie",
    "Tropen- und Reisemedizin",
    "Hebammen",
    "Chiropraktoren oder Chiropraktorinnen",
    "Aufnahme (Datum)",
    "letzte Änderung (Datum)",
)


def test_result_declares_compared_fields_and_the_workbook_columns_it_cannot_see(tmp_path, store):
    _, conn, db = store
    manifest = _imported(tmp_path, store)

    result = compare_editions(manifest, conn, db, from_date=_JAN, to_date=_JUL)

    assert result.compared_fields == (
        "category",
        "designation.de",
        "designation.fr",
        "designation.it",
        "price_chf",
        "tax_points",
        "unit",
    )
    scope = result.scope_limitation
    assert scope.columns_compared == {
        "Deutsch": ("Pos.-Nr.", "TP", "Bezeichnung"),
        "Français": ("No. Pos.", "Dénomination"),
        "Italiano": ("No. pos.", "Denominazione"),
    }
    assert scope.columns_compared_only_as_category == {
        "Deutsch": (
            "Chemie",
            "Hämatologie",
            "Immunologie",
            "Genetik",
            "Mikrobiologie",
            "Spezialanalyse",
            "Basisanalyse",
        )
    }
    assert scope.columns_not_compared["Deutsch"] == _GERMAN_COLUMNS_NOT_COMPARED
    for sheet, carried in (("Français", "Dénomination"), ("Italiano", "Denominazione")):
        assert {"Limitation", "Limitazioni"} & set(scope.columns_not_compared[sheet])
        assert carried not in scope.columns_not_compared[sheet]


def test_review_correction_is_cited_without_replacing_the_member_or_blocking(
    tmp_path, store, monkeypatch
):
    repo, conn, db = store
    monkeypatch.setenv("TARIFHUB_REVIEW_THRESHOLD", "1.01")  # flag every record for review
    if db.dialect == "postgresql":  # vector(1024) rejects the 16-dim stub (as test_review_api)
        monkeypatch.setattr("tarifhub_ingest.review_service.get_embedder", lambda *a, **k: None)
    settings = get_settings()
    january, july = _stage(tmp_path, _JAN), _stage(tmp_path, _JUL)
    import_manifest(_write_manifest(tmp_path, [january]), conn, db, settings=settings)
    decision = ReviewDecision(
        tariff_system=TariffSystem.EAL,
        tariff_code="4700",
        action="correct",
        corrections={"designation.de": "Auftragstaxe (korrigiert)"},
    )
    apply_review_decision(decision, repo=repo, audit=AuditLogger(conn, db), settings=settings)
    manifest = _write_manifest(tmp_path, [january, july])
    import_manifest(manifest, conn, db, settings=settings)

    result = compare_editions(manifest, conn, db, from_date=_JAN, to_date=_JUL)

    (position,) = [p for p in result.changed if p.tariff_code == "4700"]
    assert [c.field for c in position.changes] == ["tax_points"]  # no designation change
    assert (position.from_record.version, position.from_record.successor_versions) == (1, (2,))
    assert (position.to_record.version, position.to_record.successor_versions) == (3, ())


@pytest.mark.parametrize(
    ("from_date", "to_date", "problem"),
    [
        (_JAN, date(2026, 4, 1), "edition 2026-04-01 is not listed in the manifest"),
        (_JUL, _JAN, "2026-07-01 is not before 2026-01-01"),
        (_JAN, _JAN, "2026-01-01 is not before 2026-01-01"),
    ],
    ids=["unlisted", "reversed", "same"],
)
def test_editions_must_be_listed_and_in_date_order(tmp_path, store, from_date, to_date, problem):
    _, conn, db = store
    manifest = _write_manifest(tmp_path, [_stage(tmp_path, _JAN), _stage(tmp_path, _JUL)])

    with pytest.raises(EditionComparisonError, match=re.escape(problem)):
        compare_editions(manifest, conn, db, from_date=from_date, to_date=to_date)


def _off_hash(manifest: Path) -> Path:
    body = json.loads(manifest.read_text(encoding="utf-8"))
    body["editions"][1]["documents"][0]["sha256"] = "0" * 64
    manifest.write_text(json.dumps(body), encoding="utf-8")
    return manifest


@pytest.mark.parametrize(
    ("breakage", "problem"),
    [
        (_off_hash, "sha256 mismatch"),
        (lambda manifest: manifest.with_name("missing.json"), "manifest not found"),
    ],
    ids=["document-hash", "manifest-file"],
)
def test_manifest_and_document_problems_refuse_the_comparison(tmp_path, store, breakage, problem):
    _, conn, db = store
    manifest = breakage(_imported(tmp_path, store))

    with pytest.raises(EditionComparisonError, match=problem):
        compare_editions(manifest, conn, db, from_date=_JAN, to_date=_JUL)


def test_incomplete_stored_edition_is_refused_so_absence_is_never_guessed(tmp_path, store):
    repo, conn, db = store
    january = _stage(tmp_path, _JAN)
    path = tmp_path / january["documents"][0]["path"]
    first = next(r for r in bag_eal.parse(path) if r["tariff_code"] == "1000")
    partial = SourceSpec(
        system=TariffSystem.EAL,
        kind="bag_eal",
        path=path,
        source_url=_EAL_URL,
        valid_from=_JAN,
        source_version="BAG AL 2026-01-01",
        rows=(first,),
    )
    run_pipeline([partial], repo, AuditLogger(conn, db))  # as if an import stopped after 1000
    july_only = tmp_path / "july-only"
    import_manifest(_write_manifest(july_only, [_stage(july_only, _JUL)]), conn, db)
    manifest = _write_manifest(tmp_path, [january, _stage(tmp_path, _JUL)])

    with pytest.raises(
        EditionComparisonError,
        match=r"edition 2026-01-01: 2 of 3 position\(s\) of its full list are not stored "
        r"\(e\.g\. 1368, 4700\)",
    ):
        compare_editions(manifest, conn, db, from_date=_JAN, to_date=_JUL)


def test_stored_position_its_full_list_does_not_name_is_refused(tmp_path, store):
    repo, conn, db = store
    manifest = _imported(tmp_path, store)
    _unlisted_position(manifest, repo)

    with pytest.raises(
        EditionComparisonError,
        match=r"edition 2026-01-01: 1 stored position\(s\) of that date are not in its full "
        r"list \(e\.g\. 9999\)",
    ):
        compare_editions(manifest, conn, db, from_date=_JAN, to_date=_JUL)


def test_stored_member_with_other_content_than_its_full_list_is_refused(tmp_path, store):
    repo, conn, db = store
    seed = {
        "tariff_code": "4700",
        "designation_de": "Auftragstaxe",
        "tax_points": "99.9",
        "valid_from": _JAN,
        "source_version": "BAG AL 2026-01-01",
    }
    repo.add(freeze(map_raw(seed, system=TariffSystem.EAL, source_url=_EAL_URL)))
    manifest = _write_manifest(tmp_path, [_stage(tmp_path, _JAN), _stage(tmp_path, _JUL)])

    with pytest.raises(
        EditionComparisonError,
        match=r"edition 2026-01-01: 1 stored position\(s\) differ from its full list "
        r"\(e\.g\. 4700\)",
    ):
        compare_editions(manifest, conn, db, from_date=_JAN, to_date=_JUL)


def _undated_rows(root: Path, store) -> None:
    _, conn, db = store
    january = _stage(root, _JAN)  # the BAG file name carries no date, so the rows stay undated
    _legacy_import(conn, db, root / january["documents"][0]["path"])


def _billing_successor(root: Path, store) -> None:
    repo, _, _ = store
    _imported(root, store)
    member = next(r for r in repo.list_all() if (r.tariff_code, r.valid_from) == ("4700", _JAN))
    repo.add(freeze(member.model_copy(update={"tax_points": Decimal("99.9"), "record_hash": None})))


@pytest.mark.parametrize(
    ("history", "problem"),
    [
        (_undated_rows, "have no valid_from"),
        (_billing_successor, "changes a billing or edition field"),
    ],
    ids=["undated", "billing-successor"],
)
def test_inconsistent_stored_history_is_refused(tmp_path, store, history, problem):
    _, conn, db = store
    history(tmp_path / "setup", store)
    manifest = _write_manifest(tmp_path, [_stage(tmp_path, _JAN), _stage(tmp_path, _JUL)])

    with pytest.raises(EditionComparisonError, match=problem):
        compare_editions(manifest, conn, db, from_date=_JAN, to_date=_JUL)


def _code_1000_as_text(sheet) -> None:
    for row in sheet.iter_rows():
        if len(row) > 1 and row[1].value == 1000:
            row[1].value = "1000.00"  # the same position, written as BAG documents cite it


def test_one_position_number_written_two_ways_is_refused(tmp_path, store):
    _, conn, db = store
    july = _variant(tmp_path, _JUL, _EDITION_FIXTURES[_JUL], _code_1000_as_text)
    manifest = _write_manifest(tmp_path, [_stage(tmp_path, _JAN), july])
    import_manifest(manifest, conn, db)

    with pytest.raises(EditionComparisonError, match=r"\(e\.g\. 1000 and 1000\.00\)"):
        compare_editions(manifest, conn, db, from_date=_JAN, to_date=_JUL)


def _tp_of_4700_beyond_the_scale(sheet) -> None:
    if sheet.title == "Deutsch":
        for row in sheet.iter_rows():
            if len(row) > 2 and row[1].value == 4700:
                row[2].value = 16.60001  # five decimals; tax_points holds four


def test_value_the_canonical_scale_could_not_store_exactly_is_refused(tmp_path, store):
    _, conn, db = store
    july = _variant(tmp_path, _JUL, _EDITION_FIXTURES[_JUL], _tp_of_4700_beyond_the_scale)
    manifest = _write_manifest(tmp_path, [_stage(tmp_path, _JAN), july])
    import_manifest(manifest, conn, db)

    with pytest.raises(
        EditionComparisonError,
        match=r"edition 2026-07-01: 1 position\(s\) hold a value the canonical scale could not "
        r"store exactly \(e\.g\. 4700: tax_points\)",
    ):
        compare_editions(manifest, conn, db, from_date=_JAN, to_date=_JUL)


class _FillingAnthropic:
    """Stand-in for the ``anthropic`` client: fills the missing Italian designation."""

    def __init__(self, **_) -> None:
        self.messages = self

    def parse(self, **_):
        refinement = AIRefinement(designation_it="Traduzione generata")
        return types.SimpleNamespace(stop_reason="end_turn", parsed_output=refinement)


def _italian_of_1000_blank(sheet) -> None:
    if sheet.title == "Italiano":
        for row in sheet.iter_rows():
            if len(row) > 3 and row[1].value == 1000:
                row[3].value = None  # Denominazione: the gap the AI seam may fill


def test_ai_filled_member_value_is_refused(tmp_path, store, monkeypatch):
    _, conn, db = store
    sdk = types.ModuleType("anthropic")
    sdk.Anthropic = _FillingAnthropic
    monkeypatch.setitem(sys.modules, "anthropic", sdk)  # the SDK boundary; no network
    monkeypatch.setenv("ANTHROPIC_API_KEY", "sk-test")
    january = _variant(tmp_path, _JAN, _EDITION_FIXTURES[_JAN], _italian_of_1000_blank)
    manifest = _write_manifest(tmp_path, [january, _stage(tmp_path, _JUL)])
    import_manifest(manifest, conn, db, settings=get_settings())

    with pytest.raises(
        EditionComparisonError,
        match=r"edition 2026-01-01: 1 position\(s\) carry AI-filled fields "
        r"\(e\.g\. 1000: designation_it\)",
    ):
        compare_editions(manifest, conn, db, from_date=_JAN, to_date=_JUL)


@pytest.mark.parametrize(
    ("setup", "refused"),
    [
        (lambda manifest, repo: None, False),
        (lambda manifest, repo: _off_hash(manifest), True),  # refused before reading rows
        (_unlisted_position, True),  # refused after reading rows
    ],
    ids=["compared", "refused-on-the-documents", "refused-on-the-stored-rows"],
)
def test_comparison_writes_no_tariff_or_audit_row_even_when_refused(
    tmp_path, store, setup, refused
):
    repo, conn, db = store
    manifest = _imported(tmp_path, store)
    setup(manifest, repo)
    before = (_snapshot(repo), _audit_count(conn))

    try:
        compare_editions(manifest, conn, db, from_date=_JAN, to_date=_JUL)
    except EditionComparisonError:
        assert refused
    else:
        assert not refused

    assert (_snapshot(repo), _audit_count(conn)) == before


def _cli(*args: str) -> int:
    from tarifhub_ingest.cli import main as cli_main

    return cli_main(list(args))


def test_cli_compare_prints_the_comparison_as_json(tmp_path, store, engine, monkeypatch, capsys):
    manifest = _imported(tmp_path, store)
    monkeypatch.setenv("TARIFHUB_DB_URL", engine.db_url)

    rc = _cli("--manifest", str(manifest), "--compare", "2026-01-01", "2026-07-01")

    out, err = capsys.readouterr()
    assert (rc, err) == (0, "")
    result = json.loads(out)
    assert result["counts"] == {"added": 1, "removed": 0, "changed": 2, "unchanged": 1}
    assert [p["tariff_code"] for p in result["changed"]] == ["1368", "4700"]
    assert result["changed"][1]["changes"] == [
        {"field": "tax_points", "from_value": "21.6", "to_value": "16.6"}
    ]
    assert result["unchanged"] == ["1000"]


def test_cli_compare_output_is_byte_identical_across_runs(
    tmp_path, store, engine, monkeypatch, capsys
):
    manifest = _imported(tmp_path, store)
    monkeypatch.setenv("TARIFHUB_DB_URL", engine.db_url)

    outputs = []
    for _ in range(2):
        assert _cli("--manifest", str(manifest), "--compare", "2026-01-01", "2026-07-01") == 0
        outputs.append(capsys.readouterr().out)

    assert outputs[0] == outputs[1]


def test_cli_compare_refusal_exits_2_with_the_problems_and_writes_nothing(
    tmp_path, store, engine, monkeypatch, capsys
):
    repo, conn, _ = store
    manifest = _off_hash(_imported(tmp_path, store))
    before = (_snapshot(repo), _audit_count(conn))
    monkeypatch.setenv("TARIFHUB_DB_URL", engine.db_url)

    rc = _cli("--manifest", str(manifest), "--compare", "2026-01-01", "2026-07-01")

    out, err = capsys.readouterr()
    assert (rc, out) == (2, "")
    refusal = json.loads(err)
    assert refusal["error"] == "edition comparison refused, nothing was written"
    assert any("sha256 mismatch" in problem for problem in refusal["problems"])
    assert (_snapshot(repo), _audit_count(conn)) == before


def test_cli_compare_refuses_a_missing_database_without_creating_it(
    tmp_path, monkeypatch, capsys
):
    manifest = _write_manifest(tmp_path, [_stage(tmp_path, _JAN), _stage(tmp_path, _JUL)])
    missing = tmp_path / "missing.db"
    monkeypatch.setenv("TARIFHUB_DB_URL", f"sqlite:///{missing}")

    rc = _cli("--manifest", str(manifest), "--compare", "2026-01-01", "2026-07-01")

    assert rc == 2
    problems = json.loads(capsys.readouterr().err)["problems"]
    assert problems == ["the database cannot be opened read-only (OperationalError)"]
    assert not missing.exists()


def _database_state(engine) -> bytes | list[tuple[str, str]]:
    """What a write or a schema would change: the SQLite file's bytes, or every relation
    (table, index, view, sequence) of the Postgres database outside its system schemas."""
    if engine.label == "sqlite":
        return Path(engine.db_url.removeprefix("sqlite:///")).read_bytes()
    conn = engine.db.connect(read_only=True)
    try:
        rows = conn.execute(
            "SELECT n.nspname, c.relname FROM pg_class c"
            " JOIN pg_namespace n ON n.oid = c.relnamespace"
            " WHERE n.nspname <> 'information_schema' AND left(n.nspname, 3) <> 'pg_'"
        ).fetchall()
    finally:
        conn.close()
    return sorted((row["nspname"], row["relname"]) for row in rows)


def test_cli_compare_refuses_a_database_without_the_schema_and_leaves_it_untouched(
    tmp_path, bare_engine, monkeypatch, capsys
):
    manifest = _write_manifest(tmp_path, [_stage(tmp_path, _JAN), _stage(tmp_path, _JUL)])
    before = _database_state(bare_engine)
    assert not before  # it exists and holds nothing: an empty SQLite file, no Postgres relation
    monkeypatch.setenv("TARIFHUB_DB_URL", bare_engine.db_url)

    rc = _cli("--manifest", str(manifest), "--compare", "2026-01-01", "2026-07-01")

    out, err = capsys.readouterr()
    assert (rc, out) == (2, "")
    assert json.loads(err) == {
        "error": "edition comparison refused, nothing was written",
        "problems": [
            "the database lacks the required schema (a table the comparison reads is "
            "missing); the comparison never initialises a database"
        ],
    }
    assert _database_state(bare_engine) == before  # unchanged: no schema was initialised


@pytest.mark.parametrize(
    "args",
    [
        ("--compare", "2026-01-01", "2026-07-01"),
        ("--manifest", "m.json", "--compare", "2026-01-01", "2026-07-01", "--path", "x.xlsx"),
        ("--manifest", "m.json", "--compare", "2026-01-01", "2026-07-01", "--system", "EAL"),
        ("--manifest", "m.json", "--compare", "2026-01-01", "2026-07-01", "--refill"),
        ("--manifest", "m.json", "--compare", "2026-01-01", "2026-07-01", "--embeddings", "e5"),
        ("--manifest", "m.json", "--compare", "20260101", "2026-07-01"),
        ("--manifest", "m.json", "--compare", "2026-01-01"),
    ],
    ids=["no-manifest", "path", "system", "refill", "embeddings", "basic-date", "one-date"],
)
def test_cli_compare_usage_errors_exit_2_before_any_work(tmp_path, monkeypatch, args):
    never = tmp_path / "never.db"
    monkeypatch.setenv("TARIFHUB_DB_URL", f"sqlite:///{never}")

    with pytest.raises(SystemExit) as exit_info:
        _cli(*args)

    assert exit_info.value.code == 2
    assert not never.exists()


def test_position_absent_from_the_later_workbook_is_removed_with_its_values(tmp_path, store):
    _, conn, db = store
    later = date(2027, 1, 1)  # the January workbook, dated after the July one: 1779 drops out
    editions = [
        _stage_file(tmp_path, _JUL, _EDITION_FIXTURES[_JUL]),
        _stage_file(tmp_path, later, _EDITION_FIXTURES[_JAN]),
    ]
    manifest = _write_manifest(tmp_path, editions)
    import_manifest(manifest, conn, db)

    result = compare_editions(manifest, conn, db, from_date=_JUL, to_date=later)

    assert result.added == ()
    assert [(p.tariff_code, p.values["tax_points"]) for p in result.removed] == [("1779", "73.3")]
    assert {p.tariff_code: [c.model_dump() for c in p.changes] for p in result.changed} == {
        "1368": [{"field": "tax_points", "from_value": "64.4", "to_value": "59.4"}],
        "4700": [{"field": "tax_points", "from_value": "16.6", "to_value": "21.6"}],
    }
