"""Build the committed BAG AL edition fixtures FROM the real (gitignored) BAG downloads.

Copies real positions, unchanged, from the 1 January and 1 July 2026 BAG workbooks (all
three language sheets, banner + header layout preserved) into small fixtures that keep the
BAG download naming, so the manifest-import tests run on real edition data without the
0.8 MB originals. The ``_fixture`` pair holds four positions:

  * 1000  unchanged between the editions (76.5 TP)
  * 1368  59.4 TP in January, 64.4 TP in July
  * 4700  "Auftragstaxe", 21.6 TP in January, 16.6 TP in July
  * 1779  new in July (73.3 TP)

The ``_text_fixture`` pair (the workbook-text comparison, ADR-024) keeps 1000 and 1779 and
adds three positions whose frozen fields stay equal while German columns change:

  * 1474.1  Limitationen empty in January, a new text in July; Kumulierbarkeit
            "... 1525.00 und 1526.00" becomes "... 1525.00, 1526.00 und 1779.00"
  * 1762    "Gynäkologie und Geburtshilfe" reads "Oui" in January, "Ja" in July
  * 4707.1  Kumulierbarkeit "... maximal 18 Taxpunkte pro Tag" becomes "... 13 ...";
            its row moves down one in July, below the new 1779

Like the downloads, the fixture names carry no ``analysenliste_YYYY-MM-DD`` date.
Run from the service root:

    uv run --no-sync python tests/fixtures/eal/make_edition_fixtures.py
"""

from __future__ import annotations

import hashlib
import json
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook

from make_fixture import _POS_HEADERS, _REPO_ROOT, _SHEETS, _find_header_index

_RAW_DIR = _REPO_ROOT / "data" / "raw" / "bag" / "al"
_MANIFEST = _REPO_ROOT / "data" / "manifests" / "bag-al.json"
_POSITIONS = {1000, 1368, 1779, 4700}
_TEXT_POSITIONS = {1000, 1474.1, 1762, 1779, 4707.1}  # as the cells hold them: 1474.1 is a float
_EDITIONS = (
    (
        _RAW_DIR / "2026-01-01" / "Analysenliste per 1. Januar 2026_Excel-Format.xlsx",
        "Analysenliste per 1. Januar 2026_fixture.xlsx",
    ),
    (
        _RAW_DIR / "2026-07-01" / "Analysenliste per 1. Juli 2026_Excel-Format.xlsx",
        "Analysenliste per 1. Juli 2026_fixture.xlsx",
    ),
)
_TEXT_EDITIONS = (
    (_EDITIONS[0][0], "Analysenliste per 1. Januar 2026_text_fixture.xlsx"),
    (_EDITIONS[1][0], "Analysenliste per 1. Juli 2026_text_fixture.xlsx"),
)


def _manifest_sha256(raw: Path) -> str:
    """The SHA-256 the tracked manifest records for ``raw``."""

    manifest = json.loads(_MANIFEST.read_text(encoding="utf-8"))
    rel = raw.relative_to(_REPO_ROOT).as_posix()
    documents = [d for e in manifest["editions"] for d in e["documents"]]
    return next(d["sha256"] for d in documents if d["path"] == rel)


def build(raw: Path, fixture_name: str, positions: set[float]) -> Path:
    if not raw.exists():
        raise FileNotFoundError(
            f"raw BAG AL file not found at {raw}; download it before building the fixture"
        )
    data = raw.read_bytes()  # hash and parse the same bytes
    if hashlib.sha256(data).hexdigest() != _manifest_sha256(raw):
        raise ValueError(f"{raw} does not match its SHA-256 in {_MANIFEST}; refusing to build")

    source = load_workbook(filename=BytesIO(data), read_only=True, data_only=True)
    out = Workbook()
    out.remove(out.active)  # drop the default empty sheet

    for name in _SHEETS:
        rows = list(source[name].iter_rows(values_only=True))
        header_idx = _find_header_index(rows, _POS_HEADERS[name])
        dest = out.create_sheet(title=name)
        dest.append(list(rows[header_idx - 1]) if header_idx > 0 else [])
        dest.append(list(rows[header_idx]))
        for row in rows[header_idx + 1 :]:
            if len(row) > 1 and row[1] in positions:
                dest.append(list(row))

    source.close()
    fixture = Path(__file__).resolve().with_name(fixture_name)
    out.save(fixture)
    return fixture


if __name__ == "__main__":
    for raw, name in _EDITIONS:
        print(f"wrote fixture: {build(raw, name, _POSITIONS)}")
    for raw, name in _TEXT_EDITIONS:
        print(f"wrote fixture: {build(raw, name, _TEXT_POSITIONS)}")
